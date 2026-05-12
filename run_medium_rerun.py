"""
Re-run AI categorization for Phase_I items currently at AI_CONFIDENCE = "medium".
Splits into 200-item batches with 30s cooldown. Writes v2 results as new columns
at the right of the Phase_I_List sheet. Prints diff summary at the end.
"""
import time
import math
import sys
from datetime import datetime

import requests
import pandas as pd
from openpyxl import load_workbook

BASE_URL = "http://localhost:8000"
EXCEL_PATH = "Phase_I_PUR_2024_2025.xlsx"
SHEET_NAME = "Phase_I_List"
BATCH_SIZE = 200
COOLDOWN_SEC = 30
POLL_INTERVAL_SEC = 5
MAX_WAIT_SEC = 1200   # 20 min per batch
VECTOR_TOP_K = 10

V2_COLS = [
    "AI_MATERIAL_CATEGORY_v2",
    "AI_CONFIDENCE_v2",
    "AI_REASON_v2",
    "AI_UPDATED_AT_v2",
]


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def submit_lookup(item_numbers: list[str]) -> str:
    r = requests.post(
        f"{BASE_URL}/api/lookup/run",
        json={"item_numbers": item_numbers, "vector_top_k": VECTOR_TOP_K},
        timeout=30,
    )
    r.raise_for_status()
    return r.json()["job_id"]


def poll_job(job_id: str) -> dict:
    start = time.time()
    while True:
        r = requests.get(f"{BASE_URL}/api/batch/{job_id}/status", timeout=30)
        r.raise_for_status()
        status = r.json()
        state = status.get("status")
        if state in ("done", "error"):
            return status
        if time.time() - start > MAX_WAIT_SEC:
            status["status"] = "timeout"
            return status
        time.sleep(POLL_INTERVAL_SEC)


def main() -> int:
    # 1. Load Excel and find medium-confidence items
    log(f"Loading {EXCEL_PATH} sheet '{SHEET_NAME}'...")
    df = pd.read_excel(EXCEL_PATH, sheet_name=SHEET_NAME)
    log(f"Total rows: {len(df)}")

    medium_mask = df["AI_CONFIDENCE"].astype(str).str.lower() == "medium"
    medium_items = df.loc[medium_mask, "Material"].astype(str).tolist()
    log(f"Medium-confidence items: {len(medium_items)}")

    if not medium_items:
        log("Nothing to process.")
        return 0

    # 2. Split into batches and run
    total_batches = math.ceil(len(medium_items) / BATCH_SIZE)
    log(f"Batches: {total_batches} x {BATCH_SIZE} items, cooldown {COOLDOWN_SEC}s")

    # item_number -> new result dict
    results_by_item: dict[str, dict] = {}
    failed_batches: list[int] = []

    global_start = time.time()
    for b in range(total_batches):
        start_idx = b * BATCH_SIZE
        end_idx = min(start_idx + BATCH_SIZE, len(medium_items))
        batch = medium_items[start_idx:end_idx]
        batch_no = b + 1

        log(f"--- Batch {batch_no}/{total_batches} (items {start_idx}..{end_idx - 1}, n={len(batch)}) ---")
        batch_start = time.time()
        try:
            job_id = submit_lookup(batch)
            log(f"  job_id={job_id}")
            status = poll_job(job_id)
        except Exception as e:
            log(f"  ERROR submitting/polling: {e}")
            failed_batches.append(batch_no)
            if batch_no < total_batches:
                log(f"  Cooling down {COOLDOWN_SEC}s...")
                time.sleep(COOLDOWN_SEC)
            continue

        dur = round(time.time() - batch_start, 1)
        state = status.get("status")
        if state != "done":
            log(f"  Batch ended in state={state} after {dur}s (error/timeout)")
            failed_batches.append(batch_no)
        else:
            items = status.get("results") or []
            for row in items:
                item = row.get("Item_Number")
                if item:
                    results_by_item[item] = row
            log(
                f"  DONE in {dur}s — H={status.get('high',0)} M={status.get('medium',0)} "
                f"L={status.get('low',0)} E={status.get('error',0)} collected={len(items)}"
            )

        # Cooldown (skip on last batch)
        if batch_no < total_batches:
            log(f"  Cooling down {COOLDOWN_SEC}s...")
            time.sleep(COOLDOWN_SEC)

    total_dur_min = round((time.time() - global_start) / 60, 1)
    log(f"All batches finished in {total_dur_min} min. Collected {len(results_by_item)} results.")
    if failed_batches:
        log(f"Failed batches: {failed_batches}")

    # 3. Build v2 columns aligned to df rows
    def pick(item: str, key: str):
        r = results_by_item.get(item)
        return r.get(key, "") if r else ""

    df["AI_MATERIAL_CATEGORY_v2"] = df.apply(
        lambda row: pick(str(row["Material"]), "AI_MATERIAL_CATEGORY") if medium_mask.loc[row.name] else "",
        axis=1,
    )
    df["AI_CONFIDENCE_v2"] = df.apply(
        lambda row: pick(str(row["Material"]), "AI_confidence") if medium_mask.loc[row.name] else "",
        axis=1,
    )
    df["AI_REASON_v2"] = df.apply(
        lambda row: pick(str(row["Material"]), "AI_reason") if medium_mask.loc[row.name] else "",
        axis=1,
    )
    df["AI_UPDATED_AT_v2"] = df.apply(
        lambda row: pick(str(row["Material"]), "processed_at") if medium_mask.loc[row.name] else "",
        axis=1,
    )

    # 4. Write back to the same sheet via openpyxl (preserves other sheets)
    log(f"Writing v2 columns back to {EXCEL_PATH} / {SHEET_NAME}...")
    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    # Locate existing columns (header row 1)
    headers = {cell.value: cell.column for cell in ws[1] if cell.value}
    # Determine start column for v2 (rightmost + 1, skipping any existing v2 cols)
    max_col = ws.max_column
    start_col = max_col + 1
    # If v2 columns already exist, reuse their positions to avoid duplicates
    if all(c in headers for c in V2_COLS):
        v2_cols_idx = {c: headers[c] for c in V2_COLS}
    else:
        v2_cols_idx = {}
        for i, name in enumerate(V2_COLS):
            col = start_col + i
            ws.cell(row=1, column=col, value=name)
            v2_cols_idx[name] = col

    # Fill rows (row 2 = first data row)
    series_map = {
        "AI_MATERIAL_CATEGORY_v2": df["AI_MATERIAL_CATEGORY_v2"].tolist(),
        "AI_CONFIDENCE_v2": df["AI_CONFIDENCE_v2"].tolist(),
        "AI_REASON_v2": df["AI_REASON_v2"].tolist(),
        "AI_UPDATED_AT_v2": df["AI_UPDATED_AT_v2"].tolist(),
    }
    for col_name, values in series_map.items():
        col_idx = v2_cols_idx[col_name]
        for i, v in enumerate(values, start=2):
            ws.cell(row=i, column=col_idx, value=(v if v != "" else None))

    wb.save(EXCEL_PATH)
    log("Excel saved.")

    # 5. Diff summary: previous 'medium' → new bucket
    def normalize_conf(x):
        s = str(x).strip().lower() if x not in (None, "") else "(missing)"
        return s if s in ("high", "medium", "low", "error") else ("(missing)" if s in ("nan", "none", "") else s)

    prev_medium_df = df.loc[medium_mask].copy()
    prev_medium_df["new_conf"] = prev_medium_df["AI_CONFIDENCE_v2"].map(normalize_conf)
    prev_medium_df["prev_cat"] = prev_medium_df["AI_MATERIAL_CATEGORY"].astype(str)
    prev_medium_df["new_cat"] = prev_medium_df["AI_MATERIAL_CATEGORY_v2"].astype(str)
    prev_medium_df["cat_changed"] = prev_medium_df["prev_cat"] != prev_medium_df["new_cat"]

    log("")
    log("===== DIFF SUMMARY (previous 'medium' items) =====")
    log(f"Total previous-medium items: {len(prev_medium_df)}")
    log("New confidence distribution:")
    for conf, n in prev_medium_df["new_conf"].value_counts().items():
        pct = 100.0 * n / len(prev_medium_df)
        log(f"  {conf:<10} {n:>6}  ({pct:5.1f}%)")

    changed = int(prev_medium_df["cat_changed"].sum())
    unchanged = len(prev_medium_df) - changed
    log(f"MATERIAL_CATEGORY changed: {changed}  |  unchanged: {unchanged}")

    # Cross-tab of prev->new for visibility
    log("Transition table (prev 'medium' -> new):")
    crosstab = prev_medium_df["new_conf"].value_counts().to_dict()
    for k in ("high", "medium", "low", "error", "(missing)"):
        if k in crosstab:
            log(f"  medium -> {k:<10} {crosstab[k]}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
