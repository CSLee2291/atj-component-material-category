"""Re-categorize the 204 Phase I items where v2 confidence was low and CE has
not yet given a verdict, using BOTH Azure OpenAI and Gemini in parallel, then
write the results back into Phase_I_PUR_2024_2025_CE comfird.xlsx in-place.

Output columns (added at R..Y):
  R: AI_MATERIAL_CATEGORY_v3_azure
  S: AI_CONFIDENCE_v3_azure
  T: AI_REASON_v3_azure
  U: AI_UPDATED_AT_v3_azure
  V: AI_MATERIAL_CATEGORY_v3_gemini
  W: AI_CONFIDENCE_v3_gemini
  X: AI_REASON_v3_gemini
  Y: AI_UPDATED_AT_v3_gemini

Implementation note: openpyxl's load_workbook() on the formatted 7,968-row
file is slow (~minutes). We use pandas for the fast read+filter and only fall
back to openpyxl for the in-place write so the existing formatting / formulas
are preserved.
"""
import asyncio
import io
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import httpx
import pandas as pd
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace", line_buffering=True)
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace", line_buffering=True)

XLSX = Path(__file__).parent / "Phase_I_PUR_2024_2025_CE comfird.xlsx"
SHEET = "Phase_I_List"
API = "http://localhost:8000"
VECTOR_TOP_K = 10

NEW_HEADERS = {
    18: "AI_MATERIAL_CATEGORY_v3_azure",
    19: "AI_CONFIDENCE_v3_azure",
    20: "AI_REASON_v3_azure",
    21: "AI_UPDATED_AT_v3_azure",
    22: "AI_MATERIAL_CATEGORY_v3_gemini",
    23: "AI_CONFIDENCE_v3_gemini",
    24: "AI_REASON_v3_gemini",
    25: "AI_UPDATED_AT_v3_gemini",
}
COL_MAP = {
    "azure":  {"cat": 18, "conf": 19, "reason": 20, "updated_at": 21},
    "gemini": {"cat": 22, "conf": 23, "reason": 24, "updated_at": 25},
}


def select_target_items() -> list[str]:
    """Read the file with pandas, return Material values where AI_CONFIDENCE_v2='low'
    AND CE_MATERIAL_CATEGORY blank."""
    print("[1/4] Reading sheet via pandas (fast path)...", flush=True)
    t = time.time()
    df = pd.read_excel(XLSX, sheet_name=SHEET, engine="openpyxl")
    print(f"      pandas read in {time.time()-t:.1f}s, rows={len(df)}", flush=True)

    conf = df["AI_CONFIDENCE_v2"].astype("string").str.strip().str.lower()
    ce = df["CE_MATERIAL_CATEGORY"].astype("string").str.strip()
    mask = (conf == "low") & ((ce.isna()) | (ce == ""))
    items = df.loc[mask, "Material"].astype("string").str.strip().tolist()
    items = [i for i in items if i]
    print(f"[2/4] Selected {len(items)} target items.", flush=True)
    return items


async def submit_job(client: httpx.AsyncClient, provider: str, items: list[str]) -> str:
    r = await client.post(
        f"{API}/api/lookup/run",
        json={"item_numbers": items, "vector_top_k": VECTOR_TOP_K, "llm_provider": provider},
        timeout=60,
    )
    r.raise_for_status()
    return r.json()["job_id"]


async def poll_job(client: httpx.AsyncClient, job_id: str, label: str) -> dict:
    last_log = time.time()
    while True:
        r = await client.get(f"{API}/api/batch/{job_id}/status", timeout=30)
        r.raise_for_status()
        d = r.json()
        if d.get("status") in ("done", "error"):
            return d
        if time.time() - last_log > 15:
            print(f"        [{label}] status={d.get('status')} ...", flush=True)
            last_log = time.time()
        await asyncio.sleep(3)


async def run_both(items: list[str]) -> dict[str, list[dict]]:
    async with httpx.AsyncClient() as c:
        print(f"[3/4] Submitting {len(items)} items to /api/lookup/run for azure + gemini...", flush=True)
        azure_id, gemini_id = await asyncio.gather(
            submit_job(c, "azure", items),
            submit_job(c, "gemini", items),
        )
        print(f"      azure  job_id = {azure_id}", flush=True)
        print(f"      gemini job_id = {gemini_id}", flush=True)
        azure_job, gemini_job = await asyncio.gather(
            poll_job(c, azure_id, "azure"),
            poll_job(c, gemini_id, "gemini"),
        )

    def summarize(job, label):
        if job.get("status") != "done":
            print(f"      {label} FAILED: {job.get('error')}", flush=True)
            return []
        h = job.get("high", 0); m = job.get("medium", 0)
        l = job.get("low", 0); e = job.get("error", 0)
        t = job.get("total", 0)
        print(f"      {label} done: total={t}  high={h}  medium={m}  low={l}  error={e}", flush=True)
        return job.get("results", [])

    return {
        "azure": summarize(azure_job, "azure"),
        "gemini": summarize(gemini_job, "gemini"),
    }


def write_back(items: list[str], results: dict[str, list[dict]]) -> None:
    print("[4/4] Loading workbook with openpyxl for in-place write...", flush=True)
    t = time.time()
    wb = load_workbook(XLSX)
    print(f"      openpyxl load in {time.time()-t:.1f}s", flush=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"Sheet {SHEET!r} not found in {XLSX.name}")
    ws = wb[SHEET]

    for col, name in NEW_HEADERS.items():
        if ws.cell(1, col).value in (None, ""):
            ws.cell(1, col).value = name

    target_set = set(items)
    item_to_row: dict[str, int] = {}
    for r in range(2, ws.max_row + 1):
        m = ws.cell(r, 1).value
        if m and str(m).strip() in target_set:
            item_to_row[str(m).strip()] = r
    print(f"      mapped {len(item_to_row)} target rows", flush=True)

    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")
    written = {"azure": 0, "gemini": 0}
    for provider, recs in results.items():
        cols = COL_MAP[provider]
        for rec in recs:
            item = (rec.get("Item_Number") or "").strip()
            row = item_to_row.get(item)
            if not row:
                continue
            ws.cell(row, cols["cat"]).value = rec.get("AI_MATERIAL_CATEGORY", "")
            ws.cell(row, cols["conf"]).value = rec.get("AI_confidence", "")
            ws.cell(row, cols["reason"]).value = rec.get("AI_reason", "")
            ws.cell(row, cols["updated_at"]).value = now
            written[provider] += 1

    try:
        wb.save(XLSX)
    except PermissionError as e:
        raise SystemExit(
            f"\nCannot save {XLSX.name} — close the file in Excel and re-run. ({e})"
        )
    print(f"      Wrote: azure={written['azure']}/{len(items)}  gemini={written['gemini']}/{len(items)}", flush=True)
    print(f"      Saved {XLSX.name} in-place.", flush=True)


def main():
    items = select_target_items()
    if not items:
        print("Nothing to do.", flush=True)
        return
    if len(items) != 204:
        print(f"WARNING: expected 204 items, got {len(items)}", flush=True)

    t0 = time.time()
    results = asyncio.run(run_both(items))
    print(f"      both jobs finished in {time.time()-t0:.1f}s", flush=True)

    write_back(items, results)
    print("DONE.", flush=True)


if __name__ == "__main__":
    main()
