"""
Re-run AI categorization for Phase_I items currently at blank AI_CONFIDENCE
(skipping Material numbers that start with '142'). Splits into 200-item batches
with 30s cooldown. Writes v2 results into the existing v2 columns (reused).
Prints diff summary at the end.
"""
import time
import math
import sys
from datetime import datetime

import requests
import pandas as pd
from openpyxl import load_workbook

BASE_URL = "http://localhost:8000"
EXCEL_PATH = "Phase_I_PUR_2024_2025.xlsx"
SHEET_NAME = "Phase_I_List"
BATCH_SIZE = 200
COOLDOWN_SEC = 30
POLL_INTERVAL_SEC = 5
MAX_WAIT_SEC = 1200   # 20 min per batch
VECTOR_TOP_K = 10
SKIP_PREFIX = "142"

V2_COLS = [
    "AI_MATERIAL_CATEGORY_v2",
    "AI_CONFIDENCE_v2",
    "AI_REASON_v2",
    "AI_UPDATED_AT_v2",
]


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def submit_lookup(item_numbers: list[str]) -> str:
    r = requests.post(
        f"{BASE_URL}/api/lookup/run",
        json={"item_numbers": item_numbers, "vector_top_k": VECTOR_TOP_K},
        timeout=30,
    )
    r.raise_for_status()
    return r.json()["job_id"]


def poll_job(job_id: str) -> dict:
    start = time.time()
    while True:
        r = requests.get(f"{BASE_URL}/api/batch/{job_id}/status", timeout=30)
        r.raise_for_status()
        status = r.json()
        state = status.get("status")
        if state in ("done", "error"):
            return status
        if time.time() - start > MAX_WAIT_SEC:
            status["status"] = "timeout"
            return status
        time.sleep(POLL_INTERVAL_SEC)


def main() -> int:
    # 1. Load Excel and find blank-confidence, non-142 items
    log(f"Loading {EXCEL_PATH} sheet '{SHEET_NAME}'...")
    df = pd.read_excel(EXCEL_PATH, sheet_name=SHEET_NAME)
    log(f"Total rows: {len(df)}")

    conf_str = df["AI_CONFIDENCE"].astype(str).str.strip().str.lower()
    blank_mask = df["AI_CONFIDENCE"].isna() | (conf_str == "") | (conf_str == "nan")
    material_str = df["Material"].astype(str)
    skip_mask = material_str.str[:3] == SKIP_PREFIX
    target_mask = blank_mask & ~skip_mask

    target_items = df.loc[target_mask, "Material"].astype(str).tolist()
    log(f"Blank-confidence rows: {int(blank_mask.sum())}")
    log(f"Skipped '{SKIP_PREFIX}*' rows: {int((blank_mask & skip_mask).sum())}")
    log(f"Items to process: {len(target_items)}")

    if not target_items:
        log("Nothing to process.")
        return 0

    # 2. Split into batches and run
    total_batches = math.ceil(len(target_items) / BATCH_SIZE)
    log(f"Batches: {total_batches} x {BATCH_SIZE} items, cooldown {COOLDOWN_SEC}s")

    results_by_item: dict[str, dict] = {}
    failed_batches: list[int] = []

    global_start = time.time()
    for b in range(total_batches):
        start_idx = b * BATCH_SIZE
        end_idx = min(start_idx + BATCH_SIZE, len(target_items))
        batch = target_items[start_idx:end_idx]
        batch_no = b + 1

        log(f"--- Batch {batch_no}/{total_batches} (items {start_idx}..{end_idx - 1}, n={len(batch)}) ---")
        batch_start = time.time()
        try:
            job_id = submit_lookup(batch)
            log(f"  job_id={job_id}")
            status = poll_job(job_id)
        except Exception as e:
            log(f"  ERROR submitting/polling: {e}")
            failed_batches.append(batch_no)
            if batch_no < total_batches:
                log(f"  Cooling down {COOLDOWN_SEC}s...")
                time.sleep(COOLDOWN_SEC)
            continue

        dur = round(time.time() - batch_start, 1)
        state = status.get("status")
        if state != "done":
            log(f"  Batch ended in state={state} after {dur}s (error/timeout)")
            failed_batches.append(batch_no)
        else:
            items = status.get("results") or []
            for row in items:
                item = row.get("Item_Number")
                if item:
                    results_by_item[item] = row
            log(
                f"  DONE in {dur}s — H={status.get('high',0)} M={status.get('medium',0)} "
                f"L={status.get('low',0)} E={status.get('error',0)} collected={len(items)}"
            )

        if batch_no < total_batches:
            log(f"  Cooling down {COOLDOWN_SEC}s...")
            time.sleep(COOLDOWN_SEC)

    total_dur_min = round((time.time() - global_start) / 60, 1)
    log(f"All batches finished in {total_dur_min} min. Collected {len(results_by_item)} results.")
    if failed_batches:
        log(f"Failed batches: {failed_batches}")

    # 3. Write v2 columns back — REUSE existing v2 columns, only fill target rows
    log(f"Writing v2 columns back to {EXCEL_PATH} / {SHEET_NAME} (reusing existing columns)...")
    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    headers = {cell.value: cell.column for cell in ws[1] if cell.value}
    missing = [c for c in V2_COLS if c not in headers]
    if missing:
        max_col = ws.max_column
        for i, name in enumerate(missing):
            col = max_col + 1 + i
            ws.cell(row=1, column=col, value=name)
            headers[name] = col
        log(f"  Added missing v2 columns: {missing}")

    col_cat = headers["AI_MATERIAL_CATEGORY_v2"]
    col_conf = headers["AI_CONFIDENCE_v2"]
    col_rsn = headers["AI_REASON_v2"]
    col_ts = headers["AI_UPDATED_AT_v2"]

    # df row 0 -> worksheet row 2
    written = 0
    for i in df.index[target_mask]:
        item = str(df.at[i, "Material"])
        r = results_by_item.get(item)
        if not r:
            continue
        ws_row = i + 2
        ws.cell(row=ws_row, column=col_cat, value=r.get("AI_MATERIAL_CATEGORY") or None)
        ws.cell(row=ws_row, column=col_conf, value=r.get("AI_confidence") or None)
        ws.cell(row=ws_row, column=col_rsn, value=r.get("AI_reason") or None)
        ws.cell(row=ws_row, column=col_ts, value=r.get("processed_at") or None)
        written += 1

    wb.save(EXCEL_PATH)
    log(f"Excel saved. Cells written for {written} items.")

    # 4. Diff summary
    def normalize_conf(x):
        if x is None:
            return "(missing)"
        s = str(x).strip().lower()
        if s in ("", "nan", "none"):
            return "(missing)"
        return s

    bucket = {"high": 0, "medium": 0, "low": 0, "error": 0, "(missing)": 0, "other": 0}
    for item in target_items:
        r = results_by_item.get(item)
        if not r:
            bucket["(missing)"] += 1
            continue
        c = normalize_conf(r.get("AI_confidence"))
        if c in bucket:
            bucket[c] += 1
        else:
            bucket["other"] += 1

    log("")
    log("===== DIFF SUMMARY (previous blank-confidence, non-142 items) =====")
    log(f"Total items targeted: {len(target_items)}")
    log(f"Previous confidence: (blank) for all of these")
    log("New confidence distribution:")
    total = len(target_items)
    for k in ("high", "medium", "low", "error", "(missing)", "other"):
        n = bucket[k]
        if n == 0:
            continue
        pct = 100.0 * n / total
        log(f"  {k:<10} {n:>6}  ({pct:5.1f}%)")

    log(f"Transition: (blank) -> high={bucket['high']}, medium={bucket['medium']}, "
        f"low={bucket['low']}, error={bucket['error']}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
